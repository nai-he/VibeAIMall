from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


EA_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"
INK = "202124"
MUTED = "4B5563"
LIGHT = "6B7280"
RULE = "E5E7EB"


def apply_font(run, size: float, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def run(paragraph, text: str, *, size=9.0, color=INK, bold=False, italic=False):
    r = paragraph.add_run(text)
    apply_font(r, size, color, bold, italic)
    return r


def paragraph(doc: Document, *, before=0, after=3, line=1.08, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_bottom_rule(paragraph, color: str = RULE) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def section_heading(doc: Document, text: str) -> None:
    p = paragraph(doc, before=7, after=4, line=1.0)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=10.2, bold=True)
    add_bottom_rule(p)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.1)
    p.paragraph_format.line_spacing = 1.08
    run(p, text, size=8.75)


def skill(doc: Document, label: str, text: str) -> None:
    p = paragraph(doc, before=0, after=2.2, line=1.06)
    run(p, f"{label}：", size=8.9, bold=True)
    run(p, text, size=8.9)


def project_header(doc: Document, title: str, date: str) -> None:
    p = paragraph(doc, before=4.8, after=1.2, line=1.0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.18), WD_TAB_ALIGNMENT.RIGHT)
    run(p, title, size=9.55, bold=True)
    run(p, "\t" + date, size=8.45, color=MUTED, bold=True)


def tech(doc: Document, text: str) -> None:
    p = paragraph(doc, before=0, after=1.7, line=1.0)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=8.35, bold=True, italic=True)


def desc(doc: Document, text: str) -> None:
    p = paragraph(doc, before=0, after=2.1, line=1.06)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=8.8, color=MUTED)


def build(output: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.12)
    sec.bottom_margin = Cm(1.0)
    sec.left_margin = Cm(1.45)
    sec.right_margin = Cm(1.45)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    normal.font.size = Pt(8.9)
    normal.font.color.rgb = RGBColor.from_string(INK)

    list_style = doc.styles["List Bullet"]
    list_style.font.name = LATIN_FONT
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    list_style.font.size = Pt(8.75)
    list_style.font.color.rgb = RGBColor.from_string(INK)

    p = paragraph(doc, before=0, after=1.5, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "XXX", size=16.5, bold=True)
    run(p, "  |  AI应用开发 / Python后端 / Android AI应用（应届）", size=9.8, color=MUTED, bold=True)

    p = paragraph(doc, before=0, after=5.5, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(
        p,
        "厦门｜电话：请补充｜邮箱：13174537305@qq.com｜GitHub：https://github.com/nai-he｜期望薪资：7k-9k，可面议",
        size=8.25,
        color=MUTED,
    )

    p = paragraph(doc, before=0, after=2.5, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.18), WD_TAB_ALIGNMENT.RIGHT)
    run(p, "Xx大学  |  信息与计算科学  本科", size=8.85, color=MUTED, bold=True)
    run(p, "\t2022.09 - 2026.06", size=8.85, color=MUTED, bold=True)
    p = paragraph(doc, before=0, after=5, line=1.0)
    run(p, "全国计算机等级考试二级 Python；MathorCup 数学建模省级三等奖；iCAN 省级三等奖；传智杯程序设计省级三等奖", size=8.45, color=LIGHT)

    section_heading(doc, "求职定位")
    desc(
        doc,
        "应届 AI 应用开发方向，能从 Python/FastAPI 后端、文件解析、RAG/Agent 原型、大模型接口联调、Android/PWA 演示应用和项目文档整理做起。项目习惯采用“规则兜底 + AI 增强 + 结构化输出”，降低模型或外部接口不可用时的风险。",
    )

    section_heading(doc, "专业技能")
    skill(doc, "Python/后端", "Python、FastAPI、Jinja2、httpx、文件上传解析、接口返回、异常处理、SQLite/MySQL 基础持久化。")
    skill(doc, "大模型应用", "OpenAI/DeepSeek 兼容 API、base_url/api_key/model 配置、结构化 Prompt、结果格式化、失败兜底和人工复核。")
    skill(doc, "RAG/Agent", "做过轻量 RAG 与 Agent 流程：文档分块、TF-IDF 检索、Memory 记录、工具调用式评分、分析链路展示。")
    skill(doc, "数据/文档", "pandas、numpy、python-docx、pypdf、openpyxl，能完成 Word/PDF/Excel 解析、清洗、落库和导出。")
    skill(doc, "移动端/前端", "Android Java + XML、AndroidX、RecyclerView、SQLite、OkHttp；了解 HTML/CSS/JavaScript、React/TypeScript、PWA。")

    section_heading(doc, "项目经历")
    project_header(doc, "Job Fit Agent - 简历与岗位 JD 匹配分析系统", "2026.06")
    tech(doc, "Python / FastAPI / Jinja2 / SQLite / RAG / Agent / python-docx / pypdf / openpyxl / OpenAI API")
    desc(doc, "面向招聘筛选场景，将简历解析、岗位技能抽取、证据检索、规则评分和报告导出串成完整 Web 工具。")
    bullet(doc, "支持上传岗位 JD 与多份简历，抽取技能并输出匹配分、优势、短板、建议和候选人排名。")
    bullet(doc, "设计轻量 Agent 链路：文档读取 -> Memory 历史求职画像 -> RAG 证据检索 -> 规则评分 -> 报告生成。")
    bullet(doc, "基于技能词表、别名归一化、类别权重和缺口项生成可解释评分；无模型 Key 时仍可完成本地规则分析。")
    bullet(doc, "实现 SQLite 结果落库和 JSON、Markdown、Word、Excel 多格式导出，便于后续查询、复盘和投递材料整理。")

    project_header(doc, "Vibe智购AI - Android AI 购物助手与 PWA 展示版", "2026.06")
    tech(doc, "Android Java / XML / AndroidX / SQLite / RecyclerView / OkHttp / Chat Completions API / PWA")
    desc(doc, "面向移动端 AI 导购场景，完成 Android 原生 App 与浏览器 PWA 双形态演示。")
    bullet(doc, "Android 端覆盖登录注册、商品浏览、AI 推荐、购物车、下单、订单状态流转、评价/售后和 AI 推荐记录。")
    bullet(doc, "实现自然语言导购：识别预算、场景、品类，输出主推商品、备选方案和购买建议，支持对话式加购与跳转购物车。")
    bullet(doc, "拆分意图解析、商品匹配、推荐结果、购物车和订单服务模块；模型不可用时回退端侧规则推荐。")
    bullet(doc, "整理 PWA 版本，支持 Android/iPhone 浏览器访问演示，补足 iOS 设备无原生包时的展示路径。")

    project_header(doc, "Gold - 公募基金决策辅助分析 Agent", "2026.05 - 2026.06")
    tech(doc, "Python / FastAPI / React / TypeScript / AKShare / efinance / pandas / numpy / OpenAI API")
    desc(doc, "面向个人基金研究场景，将公开数据采集、指标计算、风险评分、回测验证和 LLM 总结封装为分析工具。")
    bullet(doc, "输入基金代码后自动获取基金基础信息、净值历史、基金画像、技术指标、风险评分和走势情景分析。")
    bullet(doc, "实现 AKShare、efinance、Tencent 多源 fallback；计算最大回撤、波动率、夏普比率、Calmar 比率、均线趋势等指标。")
    bullet(doc, "接入 LLM 生成中文分析总结，模型不可用时回退规则化文本；前端展示单基金分析、组合管理和宏观市场看板。")

    project_header(doc, "剪辑文案生成 Agent / 手语识别系统", "2026.06")
    tech(doc, "Python / JavaScript / OpenAI API / OpenCV / MediaPipe / TensorFlow-Keras / LSTM")
    desc(doc, "两个补充项目分别覆盖 AIGC 工具工作流和计算机视觉实验，作为 AI 应用开发能力补充。")
    bullet(doc, "剪辑文案 Agent 将原始文案拆成分段口播、字幕、画面建议、尾帧衔接和视频生成提示词；无 Key 时使用本地规则兜底。")
    bullet(doc, "手语识别基于 MediaPipe 21 个手部关键点与 OpenCV 摄像头流程，实现实时手势检测、手指计数和 LSTM 序列分类实验。")

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "AI应用开发求职简历-终版优化版"
    props.subject = "AI应用开发 / Python后端 / Android AI应用"
    props.comments = ""

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    build(args.out)


if __name__ == "__main__":
    main()
