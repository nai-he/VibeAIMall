from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT_EAST_ASIA = "Microsoft YaHei"
FONT_LATIN = "Calibri"
BLUE = "1F4D78"
DARK = "1F2937"
MUTED = "5B6472"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "D9E2EC"


def set_east_asia_font(run, font_name: str = FONT_EAST_ASIA) -> None:
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_spacing(paragraph, before=0, after=3, line=1.05) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_text(paragraph, text: str, *, bold=False, color=DARK, size=None):
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=4, line=1.0)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, bold=True, color=BLUE, size=11.5)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)


def add_bullet(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=2.3, line=1.05)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True, size=9.2)
        add_text(p, text[len(bold_lead):], size=9.2)
    else:
        add_text(p, text, size=9.2)


def add_project(doc: Document, title: str, role: str, tech: str, bullets: list[str]) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=1.5, line=1.0)
    p.paragraph_format.keep_with_next = True
    add_text(p, title, bold=True, color=DARK, size=10.1)
    add_text(p, f"  |  {role}", color=MUTED, size=9)
    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=2, line=1.0)
    p2.paragraph_format.keep_with_next = True
    add_text(p2, f"技术栈：{tech}", color=MUTED, size=8.8)
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_labeled_line(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2.5, line=1.05)
    add_text(p, f"{label}：", bold=True, color=BLUE, size=9.3)
    add_text(p, text, color=DARK, size=9.3)


def build_resume(output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for style_name in ("List Bullet", "List Paragraph"):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(9.2)

    # Header
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "XXX", bold=True, color=BLUE, size=20)
    add_text(p, "  AI应用开发 / Python开发 / Android AI应用（应届）", color=DARK, size=10.5)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        p,
        "厦门｜电话：请补充｜邮箱：13174537305@qq.com｜GitHub：https://github.com/nai-he｜期望薪资：7k-9k，可面议",
        color=MUTED,
        size=8.8,
    )

    # Compact positioning callout.
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.allow_autofit = False
    set_table_width(table, [7.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_borders(cell, color="D6E4F0", size="4")
    set_cell_margins(cell, top=90, bottom=90, start=140, end=140)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=1.05)
    add_text(
        p,
        "信息与计算科学本科，主攻 AI 应用开发与 Python 后端工具。能把大模型接口、规则引擎、文件解析、数据处理、SQLite 持久化和 Web/移动端展示串成完整可演示流程；熟悉“规则兜底 + AI 增强 + 结构化输出”的落地方式。",
        color=DARK,
        size=9.2,
    )

    add_section_heading(doc, "专业技能")
    add_labeled_line(
        doc,
        "AI应用",
        "OpenAI/DeepSeek 兼容 API、结构化 Prompt、RAG 检索增强、Agent 流程拆解、AI 结果兜底与复核、JSON/Markdown/Word/Excel 输出。",
    )
    add_labeled_line(
        doc,
        "Python后端",
        "Python、FastAPI、Jinja2、httpx、文件上传解析、接口联调、SQLite/MySQL 基础、python-docx、pypdf、openpyxl。",
    )
    add_labeled_line(
        doc,
        "数据与算法",
        "pandas、numpy、AKShare/efinance 数据处理、指标计算、规则评分；了解 OpenCV、MediaPipe、TensorFlow/Keras、LSTM。",
    )
    add_labeled_line(
        doc,
        "前端/移动端",
        "HTML/CSS/JavaScript、React/TypeScript 基础；Android Java + XML、AndroidX、RecyclerView、SQLite、OkHttp；了解 PWA 多端访问。",
    )
    add_labeled_line(
        doc,
        "工程工具",
        "Git、接口调试、测试用例、项目文档；使用 Cursor/Codex、Claude Code、GitHub Copilot 辅助需求拆解、编码和调试。",
    )

    add_section_heading(doc, "项目经历")
    add_project(
        doc,
        "Job Fit Agent - 简历与岗位 JD 匹配分析系统",
        "个人项目 / 主要实现",
        "Python、FastAPI、RAG、Agent、SQLite、Jinja2、python-docx、pypdf、openpyxl、OpenAI 兼容接口",
        [
            "面向招聘筛选场景，支持上传岗位 JD 与多份简历，自动抽取技能、计算匹配分，并输出优势、短板、建议和排名。",
            "实现轻量 Agent 分析链路：文档读取、历史求职画像 Memory、当前 JD/简历 RAG 证据检索、规则评分和报告生成。",
            "使用技能词表、别名归一化、类别权重和缺口项生成可解释评分；无 API Key 时仍可完成规则分析。",
            "实现 Web 批量上传、Agent 循环展示、SQLite 结果持久化，以及 JSON、Markdown、Word、Excel 多格式导出。",
        ],
    )

    add_project(
        doc,
        "Vibe智购AI - Android AI 购物助手与 PWA 展示版",
        "个人项目 / 主要实现",
        "Android Java、XML、AndroidX、SQLite、RecyclerView、OkHttp、Chat Completions API、HTML/CSS/JavaScript、PWA",
        [
            "以购物 App 为载体完成登录注册、商品浏览、购物车、下单、订单状态流转、评价/售后和 AI 推荐记录等闭环。",
            "实现 AI 导购交互：识别预算、场景、品类并输出主推商品、备选方案和购买建议；支持“就买这个”“买备选”“去购物车”等对话式加购。",
            "将 AI 意图解析、商品匹配、推荐结果、购物车和订单逻辑拆分为独立模块；真实模型不可用时回退端侧规则推荐。",
            "同步整理 PWA 版本，覆盖 Android/iPhone 浏览器访问演示，体现移动端与多端适配思路。",
        ],
    )

    add_project(
        doc,
        "Gold - 公募基金决策辅助分析 Agent",
        "个人项目 / 主要实现",
        "Python、FastAPI、React、TypeScript、AKShare、efinance、pandas/numpy、OpenAI 兼容接口",
        [
            "输入基金代码后自动获取基础信息、净值历史、基金画像、技术指标、风险评分、回测验证和走势情景分析。",
            "拆分数据获取、指标计算、风险评分、回测验证、持仓管理和 LLM 分析模块，并通过 FastAPI 提供接口。",
            "实现 AKShare、efinance、Tencent 多源数据 fallback；计算最大回撤、波动率、夏普比率、Calmar 比率、均线趋势等指标。",
            "接入 LLM 生成中文分析总结，模型不可用时回退规则化分析；前端展示单基金分析、组合管理和宏观市场看板。",
        ],
    )

    add_project(
        doc,
        "剪辑文案生成 Agent",
        "个人项目 / 补充项目",
        "Python、HTML、CSS、JavaScript、OpenAI 兼容接口",
        [
            "实现本地短视频文案工作台，根据原始文案生成分段口播脚本、字幕、画面建议、尾帧衔接说明和视频生成提示词。",
            "设计段落时长、固定段数、图片编号、尾帧前缀等配置项；无 API Key 时使用本地规则完成基础生成，配置模型后进行 AI 润色。",
        ],
    )

    add_project(
        doc,
        "手语识别系统 Sign Language Recognition",
        "个人项目 / 补充项目",
        "Python、OpenCV、MediaPipe、TensorFlow/Keras、LSTM、NumPy",
        [
            "基于 MediaPipe 21 个手部关键点和 OpenCV 摄像头流程，实现实时手势检测、手指计数和自定义手语动作识别实验。",
            "完成数据采集、视频切分、关键点提取、LSTM 训练和实时识别脚本；使用连续 30 帧序列和置信度阈值过滤低可信结果。",
        ],
    )

    add_section_heading(doc, "教育背景与证书")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.05)
    add_text(p, "Xx大学  |  信息与计算科学  |  本科  |  2022.09 - 2026.06", bold=True, color=DARK, size=9.4)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.05)
    add_text(p, "主修课程：高等数学、线性代数、概率论与数理统计、数据结构、算法设计、Python 程序设计、数据库基础、计算方法。", size=9.2)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.05)
    add_text(
        p,
        "证书/竞赛：全国计算机等级考试二级 Python；MathorCup 数学建模省级三等奖；iCAN 省级三等奖；数学能力挑战赛一等奖；传智杯程序设计省级三等奖。",
        size=9.2,
    )

    add_section_heading(doc, "可胜任工作")
    add_bullet(doc, "Python/FastAPI 接口开发、文件上传解析、第三方 API 接入、AI 分析结果落库与多格式导出。")
    add_bullet(doc, "大模型应用中的 Prompt 调整、接口调试、异常兜底、结果格式化、前后端联调和基础页面展示。")
    add_bullet(doc, "Android/PWA 演示型应用功能迭代、业务流程梳理、项目文档整理和 GitHub 求职作品集维护。")

    # Lightweight footer.
    for i, sec in enumerate(doc.sections):
        footer = sec.footer.paragraphs[0]
        set_paragraph_spacing(footer, before=0, after=0, line=1.0)
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(footer, "AI应用开发求职简历｜项目可通过 GitHub/本地演示进一步说明", color="8A94A6", size=8)

    # Scrub common metadata.
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "AI应用开发求职简历"
    props.subject = "AI应用开发 / Python开发 / Android AI应用"
    props.comments = ""

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    build_resume(args.out)


if __name__ == "__main__":
    main()
