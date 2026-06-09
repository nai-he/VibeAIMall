from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


EA_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"
INK = "222222"
MUTED = "555555"
LIGHT = "888888"


def font(run, size: float, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def para(doc: Document, before=0, after=3, line=1.08, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align:
        p.alignment = align
    return p


def add_run(p, text: str, size=9, color=INK, bold=False, italic=False):
    r = p.add_run(text)
    font(r, size=size, color=color, bold=bold, italic=italic)
    return r


def heading(doc: Document, text: str) -> None:
    p = para(doc, before=8, after=4, line=1.0)
    p.paragraph_format.keep_with_next = True
    add_run(p, text, size=10.5, bold=True)


def bullet(doc: Document, text: str, size=8.9) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.08
    add_run(p, text, size=size, color=INK)


def title_line(doc: Document, title: str, right: str) -> None:
    p = para(doc, before=5, after=1.5, line=1.0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.18), WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, title, size=9.6, bold=True)
    add_run(p, "\t" + right, size=8.7, color=MUTED, bold=True)


def tech_line(doc: Document, text: str) -> None:
    p = para(doc, before=0, after=2, line=1.0)
    p.paragraph_format.keep_with_next = True
    add_run(p, text, size=8.55, color=INK, bold=True, italic=True)


def build(output: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.15)
    sec.bottom_margin = Cm(1.05)
    sec.left_margin = Cm(1.45)
    sec.right_margin = Cm(1.45)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    normal.font.size = Pt(9)

    list_style = doc.styles["List Bullet"]
    list_style.font.name = LATIN_FONT
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    list_style.font.size = Pt(8.9)

    p = para(doc, before=0, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, "XXX", size=17, bold=True)
    add_run(p, "  |  AI应用开发 / Python开发 / Android AI应用（应届）", size=10, color=MUTED, bold=True)

    p = para(doc, before=0, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(
        p,
        "厦门｜电话：请补充｜邮箱：13174537305@qq.com｜GitHub：https://github.com/nai-he｜期望薪资：7k-9k，可面议",
        size=8.3,
        color=MUTED,
    )

    p = para(doc, before=0, after=7, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.18), WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, "Xx大学  |  信息与计算科学  本科  |  全国计算机二级 Python", size=8.9, color=MUTED, bold=True)
    add_run(p, "\t2022.09 - 2026.06", size=8.9, color=MUTED, bold=True)

    heading(doc, "专业技能")
    bullet(doc, "熟悉 Python 语言与 FastAPI 接口开发，能完成文件上传、接口返回、前后端联调、异常处理与 SQLite/MySQL 基础持久化。")
    bullet(doc, "掌握大模型应用基础流程：OpenAI/DeepSeek 兼容 API、结构化 Prompt、RAG 检索增强、Agent 流程拆解、AI 结果兜底与复核。")
    bullet(doc, "熟悉 Word/PDF/Excel 文档解析与导出，使用过 python-docx、pypdf、openpyxl、pandas、numpy 完成数据清洗和结构化处理。")
    bullet(doc, "有 Android Java + XML 项目经验，使用 AndroidX、RecyclerView、SQLite、OkHttp 完成 AI 导购、购物车、订单等移动端业务。")
    bullet(doc, "了解 HTML/CSS/JavaScript、React/TypeScript、OpenCV、MediaPipe、TensorFlow/Keras；能使用 Codex、Claude Code、Cursor 等 AI Coding 工具辅助迭代。")

    heading(doc, "项目经历")
    title_line(doc, "Job Fit Agent - 简历与岗位 JD 匹配分析系统", "2026.06")
    tech_line(doc, "Python / FastAPI / Jinja2 / SQLite / RAG / Agent / python-docx / pypdf / openpyxl / OpenAI API")
    bullet(doc, "面向招聘筛选场景，支持上传岗位 JD 与多份简历，自动抽取技能、计算匹配分，并输出优势、短板、建议和排名报告。")
    bullet(doc, "设计轻量 Agent 链路：文档读取 -> Memory 历史求职画像 -> RAG 证据检索 -> 规则评分 -> 报告生成，页面可展开分析过程。")
    bullet(doc, "基于技能词表、别名归一化、类别权重和缺口项生成可解释评分；无模型 Key 时仍可完成本地规则分析。")
    bullet(doc, "实现 Web 批量上传、SQLite 结果落库，以及 JSON、Markdown、Word、Excel 多格式导出，方便后续查询和复盘。")

    title_line(doc, "Vibe智购AI - Android AI 购物助手与 PWA 展示版", "2026.06")
    tech_line(doc, "Android Java / XML / AndroidX / SQLite / RecyclerView / OkHttp / Chat Completions API / PWA")
    bullet(doc, "以购物 App 为载体，完成登录注册、商品浏览、AI 推荐、购物车、下单、订单状态流转、评价/售后和 AI 推荐记录闭环。")
    bullet(doc, "实现 AI 导购：识别用户预算、场景、品类并输出主推商品、备选方案和购买建议，支持“就买这个”“买备选”“去购物车”等指令。")
    bullet(doc, "拆分意图解析、商品匹配、推荐结果、购物车和订单服务模块；真实模型不可用时回退端侧规则推荐，保证演示链路可运行。")
    bullet(doc, "同步整理 PWA 版本，支持 Android/iPhone 浏览器访问演示，体现移动端业务与多端适配思路。")

    title_line(doc, "Gold - 公募基金决策辅助分析 Agent", "2026.05 - 2026.06")
    tech_line(doc, "Python / FastAPI / React / TypeScript / AKShare / efinance / pandas / numpy / OpenAI API")
    bullet(doc, "输入基金代码后自动获取基础信息、净值历史、基金画像、技术指标、风险评分、回测验证和走势情景分析。")
    bullet(doc, "拆分数据获取、指标计算、风险评分、回测验证、持仓管理和 LLM 分析模块，并通过 FastAPI 提供接口。")
    bullet(doc, "实现 AKShare、efinance、Tencent 多源 fallback；计算最大回撤、波动率、夏普比率、Calmar 比率、均线趋势等指标。")
    bullet(doc, "接入 LLM 生成中文分析总结，模型不可用时回退规则化分析；前端展示单基金分析、组合管理和宏观市场看板。")

    title_line(doc, "剪辑文案生成 Agent / 手语识别系统", "2026.06")
    tech_line(doc, "Python / HTML / CSS / JavaScript / OpenAI API / OpenCV / MediaPipe / TensorFlow-Keras / LSTM")
    bullet(doc, "剪辑文案 Agent 支持将原始文案拆成分段口播、字幕、画面建议、尾帧衔接和视频生成提示词；无 Key 时使用本地规则兜底。")
    bullet(doc, "手语识别项目基于 MediaPipe 21 个手部关键点与 OpenCV 摄像头流程，实现实时手势检测、手指计数和 LSTM 序列分类实验。")

    heading(doc, "证书与竞赛")
    bullet(doc, "全国计算机等级考试二级 Python；MathorCup 数学建模省级三等奖；iCAN 省级三等奖；数学能力挑战赛一等奖；传智杯程序设计省级三等奖。")

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "AI应用开发求职简历-极简技术岗版"
    props.subject = "AI应用开发 / Python开发 / Android AI应用"
    props.comments = ""

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    build(args.out)


if __name__ == "__main__":
    main()
